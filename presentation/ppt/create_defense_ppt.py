from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
ASSET = ROOT.parent / "latex" / "assets" / "ug_template_cover_bg.jpg"
OUTPUT = ROOT / "defense_slides.pptx"

BLUE = RGBColor(21, 55, 94)
WHITE = RGBColor(255, 255, 255)


def add_text(slide, text, left, top, width, height, size=24, bold=False,
             color=BLUE, align=PP_ALIGN.LEFT, font="Aptos"):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_bullets(slide, items, left=0.9, top=1.55, width=11.5, height=5.3,
                size=21):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.08)
    frame.margin_right = Inches(0.04)
    for index, item in enumerate(items):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = item
        paragraph.level = 0
        paragraph.font.name = "Aptos"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = BLUE
        paragraph.space_after = Pt(12)
    return box


def add_title(slide, title):
    slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.78)
    ).fill.solid()
    bar = slide.shapes[-1]
    bar.fill.fore_color.rgb = BLUE
    bar.line.fill.background()
    add_text(slide, title, 0.55, 0.08, 12.2, 0.6, size=26, bold=True,
             color=WHITE)


def cover(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_picture(str(ASSET), 0, 0, width=Inches(13.333), height=Inches(7.5))
    add_text(slide, "FINAL YEAR PROJECT PRESENTATION & DEMO",
             1.0, 0.25, 11.3, 0.35, size=18, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER)
    add_text(slide, "Resource-Constrained Evaluation of Lightweight\n"
             "Concept Drift Detectors in Machine Learning",
             1.35, 0.85, 10.65, 0.75, size=18, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER)
    add_text(slide, "Stephen Kwaku Azongo\nID: 11027858\n"
             "Department of Computer Science, University of Ghana",
             1.1, 1.58, 11.1, 0.9, size=15, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER)
    add_text(slide, "SUPERVISOR: KOFI SARPONG ADU-MANU\n"
             "8TH OCTOBER 2026",
             1.5, 2.55, 10.3, 0.48, size=11, bold=True, color=WHITE,
             align=PP_ALIGN.CENTER)


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    cover(prs)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "OUTLINE OF PRESENTATION")
    add_bullets(slide, [
        "1. Introduction",
        "2. Aims and Objectives",
        "3. Related Work",
        "4. Proposed Solution (Methodology)",
        "5. Key Results",
        "6. Conclusion",
    ], top=1.45, size=23)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Introduction")
    add_bullets(slide, [
        "Problem context: Production machine-learning models face concept drift as P(X, y) changes over time.",
        "Existing benchmarks usually hold deployment resources constant.",
        "Real systems must trade off monitoring frequency, detection-window size, and reference memory.",
        "Study focus: evaluate lightweight unsupervised detectors across financial, telecom, and industrial datasets.",
    ])

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Aims and Objectives")
    add_bullets(slide, [
        "Aim: determine whether monitoring configuration is more consequential than detector algorithm.",
        "Compare KS, PSI, and Wasserstein detectors over 8,100 experimental runs.",
        "Quantify the effects of monitoring frequency, window size, and reference size.",
        "Evaluate sudden, gradual, and recurring drift.",
        "Assess Random Forest, Logistic Regression, and LightGBM robustness.",
    ], size=19)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Related Work")
    add_bullets(slide, [
        "Concept-drift foundations: drift types and adaptive monitoring.",
        "Empirical benchmarks compare detectors but rarely model deployment constraints as primary factors.",
        "KS, PSI, and Wasserstein provide practical, label-free, lightweight monitoring.",
        "Research gap: no systematic ranking of resource constraints by quantitative impact on detection outcomes.",
    ], size=20)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Proposed Solution (Methodology)")
    add_bullets(slide, [
        "Datasets: UCI Credit Default, IBM Telco Churn, and UCI AI4I Predictive Maintenance.",
        "Drift simulation: sudden, gradual, and recurring drift after a 60/40 reference-stream split.",
        "Detectors: Kolmogorov-Smirnov, Population Stability Index, and Wasserstein distance.",
        "Constraint grid: frequency {10, 50, 100, 500, 1000}; window {50, 100, 500, 1000}; reference {100, 500, unlimited}.",
        "Analysis: Kruskal-Wallis, Mann-Whitney U, Bonferroni correction, Spearman correlation, and effect sizes.",
    ], size=17)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Key Results")
    add_bullets(slide, [
        "Constraint configuration produced a 37-fold range in detection delay.",
        "Gradual drift was significantly harder to detect than sudden or recurring drift (r ≈ 0.30).",
        "Detection-window size had the strongest individual impact; monitoring frequency was second.",
        "A 500-sample reference window performed comparably to unlimited reference in most settings.",
        "Random Forest degraded less than LightGBM (r = 0.383) and Logistic Regression (r = 0.185).",
    ], size=18)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "Conclusion")
    add_bullets(slide, [
        "Main conclusion: configuration matters more than algorithm for practical drift-detection performance.",
        "1. Configure the monitoring regime first.",
        "2. Account for drift type, especially gradual drift.",
        "3. Then optimize detector and predictive-model selection.",
    ], size=21)
    add_text(slide, "THANK YOU", 4.4, 5.65, 4.5, 0.55, size=28, bold=True,
             color=BLUE, align=PP_ALIGN.CENTER)

    prs.save(OUTPUT)


if __name__ == "__main__":
    build()
